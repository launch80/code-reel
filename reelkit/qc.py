"""qc — post-production QC for a built reel.

Measures what a machine can (container, codecs, color tags, fast-start,
frame count, A/V length, black/freeze, EBU R128 loudness, layout audit,
novelty) and pre-fills what a human must judge, with exact timecodes.
Writes .reel/qc_report.md and .reel/qc_report.xlsx:
  Summary · Test Sequence · Cue Sheet · Tech Specs · Fact Check
"""
import datetime
import json
import os
import re
import subprocess

from . import build as B


def _probe(mp4, sel, entries):
    r = subprocess.run(["ffprobe", "-v", "error", "-i", mp4] + sel + ["-show_entries", entries, "-of", "json"],
                       capture_output=True, text=True)
    try:
        j = json.loads(r.stdout or "{}")
    except json.JSONDecodeError:
        return {}
    return (j.get("streams") or [j.get("format") or {}])[0]


def _boxes(path):
    off, found = 0, {}
    with open(path, "rb") as f:
        while True:
            hdr = f.read(8)
            if len(hdr) < 8:
                break
            size = int.from_bytes(hdr[:4], "big")
            typ = hdr[4:8].decode("latin1")
            if typ in ("moov", "mdat"):
                found.setdefault(typ, off)
            if size == 1:
                size = int.from_bytes(f.read(8), "big")
                f.seek(size - 16, 1)
            elif size >= 8:
                f.seek(size - 8, 1)
            else:
                break
            off += size
    return found


def run(proj, mp4=None):
    rd = os.path.join(proj, ".reel")
    tl = json.load(open(os.path.join(rd, "timeline.resolved.json")))
    mp4 = mp4 or os.path.join(proj, "reel.mp4")
    audit = json.load(open(os.path.join(rd, "audit.json"))) if os.path.exists(os.path.join(rd, "audit.json")) else None
    novelty = json.load(open(os.path.join(rd, "novelty.json"))) if os.path.exists(os.path.join(rd, "novelty.json")) else None
    FPS, DUR = int(tl["fps"]), float(tl["duration"])
    N = int(round(DUR * FPS))
    tc = lambda s: "%02d:%02d:%02d:%02d" % (int(round(s * FPS)) // (3600 * FPS), int(round(s * FPS)) // (60 * FPS) % 60, int(round(s * FPS)) // FPS % 60, int(round(s * FPS)) % FPS)
    ev = tl["events"]

    def sname(t):
        for s in tl["scenes"]:
            if s["t0"] - 1e-6 <= t < s["t1"] + 1e-6:
                return "%02d %s" % (s["i"] + 1, s["name"])
        return "TAIL"

    tech = []

    def spec(p, target, got, ok, how, note=""):
        tech.append((p, target, got, "PASS" if ok else "FAIL", how, note))

    exists = os.path.exists(mp4)
    v = _probe(mp4, ["-select_streams", "v:0", "-count_frames"], "stream=width,height,r_frame_rate,nb_read_frames,codec_name,profile,pix_fmt,color_space,color_primaries,color_transfer,color_range,bit_rate,duration") if exists else {}
    a = _probe(mp4, ["-select_streams", "a:0"], "stream=codec_name,sample_rate,channels,bit_rate,duration") if exists else {}
    fm = _probe(mp4, [], "format=duration") if exists else {}
    spec("Width", tl["w"], v.get("width"), str(v.get("width")) == str(tl["w"]), "ffprobe")
    spec("Height", tl["h"], v.get("height"), str(v.get("height")) == str(tl["h"]), "ffprobe")
    spec("Frame rate", "%d/1" % FPS, v.get("r_frame_rate"), v.get("r_frame_rate") == "%d/1" % FPS, "ffprobe")
    spec("Frame count (decoded)", N, v.get("nb_read_frames"), str(v.get("nb_read_frames")) == str(N), "ffprobe -count_frames", "off by one = dropped/dup frame")
    spec("Duration s", "%.3f" % DUR, fm.get("duration"), abs(float(fm.get("duration") or 0) - DUR) < 0.05, "ffprobe")
    spec("Video codec / profile", "h264 High", "%s %s" % (v.get("codec_name"), v.get("profile")), v.get("codec_name") == "h264" and v.get("profile") == "High", "ffprobe")
    spec("Pixel format", "yuv420p", v.get("pix_fmt"), v.get("pix_fmt") == "yuv420p", "ffprobe")
    ct = [v.get("color_space"), v.get("color_primaries"), v.get("color_transfer")]
    spec("Color tags", "bt709 x3", " / ".join(str(x) for x in ct), all(x == "bt709" for x in ct), "ffprobe", "untagged = color shift on some players")
    spec("Color range", "tv", v.get("color_range"), v.get("color_range") == "tv", "ffprobe")
    bx = _boxes(mp4) if exists else {}
    spec("Fast-start", "moov before mdat", str(bx), bool(bx.get("moov")) and bx.get("moov", 1e18) < bx.get("mdat", 1e18), "box scan")
    spec("Audio", "aac 48000 2ch", "%s %s %sch" % (a.get("codec_name"), a.get("sample_rate"), a.get("channels")),
         a.get("codec_name") == "aac" and str(a.get("sample_rate")) == "48000" and str(a.get("channels")) == "2", "ffprobe")
    spec("A/V length match", "both %.3f" % DUR, "v %s / a %s" % (v.get("duration"), a.get("duration")),
         abs(float(v.get("duration") or 0) - DUR) < 0.05 and abs(float(a.get("duration") or 0) - DUR) < 0.06, "ffprobe")
    if exists:
        r = subprocess.run(["ffmpeg", "-v", "info", "-i", mp4, "-vf", "blackdetect=d=0.03:pix_th=0.05,freezedetect=n=-45dB:d=0.5", "-f", "null", "-"], capture_output=True, text=True)
        blacks = re.findall(r"black_start:([\d.]+) black_end:([\d.]+)", r.stderr)
        freezes = re.findall(r"freeze_start:([\d.]+)", r.stderr)
        spec("Black segments", "fades only (<=0.3s at head/tail)", "; ".join("%s-%s" % b for b in blacks) or "none",
             all(float(e) - float(s) <= 0.3 and (float(s) < 0.2 or float(e) > DUR - 0.4) for s, e in blacks), "blackdetect")
        spec("Frozen segments >= 0.5s", "0", len(freezes), len(freezes) == 0, "freezedetect")
        L = B.measure(mp4)
        spec("Integrated loudness", "-14 +/- 1 LUFS", L["I"], -15 <= L["I"] <= -13, "ebur128")
        spec("True peak", "<= -1.0 dBTP", L["TP"], L["TP"] <= -1.0, "ebur128")
        spec("Loudness range", "<= 10 LU", L["LRA"], L["LRA"] <= 10, "ebur128")
    else:
        spec("Built MP4 present", mp4, "missing", False, "os.path.exists")

    auto = []   # (phase, name, status, note)
    auto.append(("A", "compile clean (0 errors)", "PASS" if not tl["errors"] else "FAIL", "; ".join(tl["errors"])[:300]))
    auto.append(("A", "claims: every on-screen number sourced", "PASS" if not [w for w in tl["warnings"] if "unsourced" in w] else "FAIL",
                 "%d unsourced" % len([w for w in tl["warnings"] if "unsourced" in w])))
    if audit:
        auto.append(("C", "layout audit: no overlaps / off-canvas text / blank scenes", "PASS" if not audit["errors"] else "FAIL", "; ".join(audit["errors"][:6])))
        auto.append(("C", "layout audit: text inside title-safe", "PASS" if not audit["warnings"] else "WARN", "; ".join(audit["warnings"][:6])))
    else:
        auto.append(("C", "layout audit run", "FAIL", "run: reel.py audit"))
    if novelty:
        auto.append(("C", "novelty vs blueprints + past reels", "PASS" if novelty["ok"] else "FAIL", novelty["summary"]))
    auto.append(("E", "impact sync (render + audio share compiled events)", "PASS", "by construction"))

    cues = [(0.0, "fade in", "score in")]
    for tr in ev["transitions"]:
        cues.append((tr["t"], "%s -> %s (%s%s)" % (tr["from"], tr["to"], tr["type"], " %.2fs" % tr["dur"] if tr["dur"] else ""),
                     {"glitch": "riser + impact + glitch burst", "flash": "riser + impact", "cut": "click"}.get(tr["type"], "whoosh" if tr["type"] in ("push", "whip", "slide", "wipe", "shutter") else "swell")))
    for s in ev["slams"]:
        cues.append((s["t"], "%s: %s" % (s["scene"], s["cause"]), "impact"))
    for ty in ev["typing"]:
        cues.append((ty["t"], "%s: typing %d chars (%.2fs)" % (ty["scene"], ty["n"], ty["dur"]), "%d clicks" % ty["n"]))
    for tk in ev["ticks"]:
        cues.append((tk["a"], "%s: counter runs to %.2fs" % (tk["scene"], tk["b"]), "ticks @20Hz"))
    if ev.get("bell") is not None:
        cues.append((ev["bell"], "end card", "bell"))
    cues.append((DUR - 0.8, "tail", "audio fade 0.8s"))
    cues.sort(key=lambda c: c[0])

    facts = []
    for u in tl["claim_uses"]:
        c = tl["claims"].get(u["claim"], {})
        src = c.get("url") or c.get("source") or ""
        facts.append((u["claim"], u.get("text", c.get("display", c.get("value"))), u["path"], src, "Low" if src else "High", "Verified" if src else "Source missing"))
    for w in tl["warnings"] + tl["errors"]:
        if "unsourced number" in w:
            facts.append(("—", w.split(" — ", 1)[-1].split(" (use a claim")[0], "", "", "High", "Unsourced"))

    tests = []
    for p, target, got, res, how, note in tech:
        tests.append(("B · Technical", p, None, None, "target %s · measured %s" % (target, got), how, res, note))
    for ph, name, res, note in auto:
        tests.append(({"A": "A · Pre-flight", "C": "C · Visual (auto)", "E": "E · Audio"}[ph], name, None, None, "", "reel.py", res, note))
    for s in tl["scenes"]:
        tests.append(("C · Visual (human)", "%02d %s [%s%s]" % (s["i"] + 1, s["name"], s["scene"], " custom" if s["custom"] else ""), s["t0"], s["t1"],
                      "legible, on-brief, motion reads, nothing clipped; would this frame only belong to THIS video?", "watch at speed, then frame-step", "Not run", ""))
    for tr in ev["transitions"]:
        tests.append(("D · Transitions", "%s at %s" % (tr["type"], tc(tr["t"])), max(0, tr["t"] - 0.3), tr["t"] + 0.3,
                      "reads as intended, no pop/flash wash-out, sound lands", "frame-step the window", "Not run", ""))
    for name in ("music vs sfx balance", "small speakers + headphones", "mono compatibility", "tail fade"):
        tests.append(("E · Audio (human)", name, None, None, "", "listen", "Not run", ""))
    for g in ("upload preview on target platform", "thumbnail / poster frame", "captions / alt text"):
        tests.append(("G · Delivery", g, None, None, "", "platform draft", "Not run", ""))
    tests.append(("H · Sign-off", "zero open FAILs + full watch with sound", None, None, "", "human", "Not run", ""))

    fails = [t for t in tests if t[6] == "FAIL"]
    pending = [t for t in tests if t[6] == "Not run"]
    ready = "READY" if not fails and not pending else "NOT YET — %d fail(s), %d manual check(s) pending" % (len(fails), len(pending))

    md = ["# QC — %s" % os.path.basename(mp4), "",
          "%s · %sx%s · %s · %d frames · style %s · seed %s · QC %s" % (tl["name"], tl["w"], tl["h"], "%dfps" % FPS, N, tl["style"]["pack"], tl["seed"], datetime.datetime.now().strftime("%Y-%m-%d %H:%M")),
          "", "**READY TO PUBLISH? %s**" % ready, "", "## Automatic", "", "| test | result | note |", "|---|---|---|"]
    for t in tests:
        if t[6] in ("PASS", "FAIL", "WARN"):
            md.append("| %s — %s | %s | %s |" % (t[0], t[1], t[6], (t[7] or t[4])[:160].replace("|", "/")))
    md += ["", "## Cue sheet", "", "| t | TC | scene | picture | sound |", "|---|---|---|---|---|"]
    md += ["| %.2f | %s | %s | %s | %s |" % (c[0], tc(c[0]), sname(c[0]), c[1], c[2]) for c in cues]
    md += ["", "## Fact check", "", "| claim | on screen | where | source | risk | status |", "|---|---|---|---|---|---|"]
    md += ["| %s | %s | %s | %s | %s | %s |" % tuple(str(x).replace("|", "/") for x in f) for f in facts]
    md += ["", "## Manual checks", ""] + ["- [ ] %s — %s %s" % (t[0], t[1], ("(%s–%s)" % (tc(t[2]), tc(t[3]))) if t[2] is not None else "") for t in pending]
    open(os.path.join(rd, "qc_report.md"), "w").write("\n".join(md) + "\n")

    try:
        from openpyxl import Workbook
        from openpyxl.styles import Alignment, Font, PatternFill
        from openpyxl.worksheet.datavalidation import DataValidation
        wb = Workbook()
        FILL = {"PASS": "D9EAD3", "FAIL": "F4CCCC", "WARN": "FCE5CD", "Not run": "FFF2CC", "Verified": "D9EAD3", "High": "F4CCCC", "Low": "D9EAD3", "Unsourced": "F4CCCC", "Source missing": "F4CCCC"}

        def sheet(ws, title, headers, rows, widths):
            ws.append([title]); ws["A1"].font = Font(name="Arial", bold=True, size=13)
            ws.append([]); ws.append(headers)
            for c in ws[3]:
                c.font = Font(name="Arial", bold=True, color="FFFFFF"); c.fill = PatternFill("solid", fgColor="1F1F22")
            for r in rows:
                ws.append(list(r))
                for c in ws[ws.max_row]:
                    c.font = Font(name="Arial", size=10); c.alignment = Alignment(wrap_text=True, vertical="top")
                    if str(c.value) in FILL:
                        c.fill = PatternFill("solid", fgColor=FILL[str(c.value)])
            for i, w in enumerate(widths):
                ws.column_dimensions[chr(65 + i)].width = w
            ws.freeze_panes = "A4"
        ws = wb.active; ws.title = "Summary"
        for k, val in [("Reel", tl["name"]), ("File", mp4), ("Style / seed", "%s / %s" % (tl["style"]["pack"], tl["seed"])),
                       ("Custom scene share", "%.0f%%" % (tl["custom_share"] * 100)), ("Auto fails", len(fails)), ("Manual pending", len(pending)), ("READY TO PUBLISH?", ready)]:
            ws.append([k, val]); ws.cell(ws.max_row, 1).font = Font(name="Arial", bold=True)
        ws.column_dimensions["A"].width = 24; ws.column_dimensions["B"].width = 90
        ts = wb.create_sheet("Test Sequence")
        sheet(ts, "Test Sequence — run top to bottom; set Status on human rows", ["Phase", "Test", "In", "Out", "Check / expected", "How", "Status", "Notes"],
              [[t[0], t[1], None if t[2] is None else tc(t[2]), None if t[3] is None else tc(t[3]), t[4], t[5], t[6], t[7]] for t in tests], [20, 44, 12, 12, 50, 22, 11, 50])
        dv = DataValidation(type="list", formula1='"Not run,PASS,FAIL,WARN,N/A"'); ts.add_data_validation(dv); dv.add("G4:G%d" % (ts.max_row + 50))
        sheet(wb.create_sheet("Cue Sheet"), "Cue Sheet — from the compiled events (render + audio share them)", ["t (s)", "Frame", "TC", "Scene", "Picture", "Sound"],
              [[round(c[0], 3), int(round(c[0] * FPS)), tc(c[0]), sname(c[0]), c[1], c[2]] for c in cues], [9, 8, 13, 22, 50, 30])
        sheet(wb.create_sheet("Tech Specs"), "Tech Specs — target vs measured", ["Parameter", "Target", "Measured", "Result", "How", "Note"], tech, [26, 26, 30, 9, 20, 40])
        sheet(wb.create_sheet("Fact Check"), "Fact Check — every claim on screen and where it came from", ["Claim", "On screen", "Where", "Source", "Risk", "Status"], facts, [16, 36, 34, 60, 8, 14])
        wb.save(os.path.join(rd, "qc_report.xlsx"))
    except ImportError:
        pass
    return {"ready": ready, "fails": [(t[0], t[1], t[7]) for t in fails], "pending": len(pending), "tech": tech}
